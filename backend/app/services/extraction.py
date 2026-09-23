"""
Step 5: text extraction.

- PDF: PyMuPDF (fitz) pulls text per page. Any page with little/no
  extractable text (i.e. a scanned image page) is rendered to an image
  and run through pytesseract OCR instead.
- DOCX: python-docx pulls paragraph text directly (DOCX is never
  "scanned" — it's always native text — so no OCR path needed).

Raises ExtractionError on failure so the caller can mark the document
"failed" with a clear message instead of crashing the background task.
"""

import io

import fitz  # PyMuPDF
import pytesseract
from PIL import Image
from docx import Document as DocxDocument

# Below this many characters, treat a PDF page as "no native text" and OCR it.
MIN_CHARS_FOR_NATIVE_TEXT = 20

# Render scale for OCR — higher = more accurate but slower. 2x is a good
# balance for typical scanned contracts.
OCR_RENDER_ZOOM = 2.0


class ExtractionError(Exception):
    pass


def extract_text(file_path: str, content_type: str) -> str:
    if content_type == "application/pdf":
        return _extract_pdf(file_path)
    if content_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        return _extract_docx(file_path)
    raise ExtractionError(f"Unsupported content type for extraction: {content_type}")


def _extract_pdf(file_path: str) -> str:
    try:
        doc = fitz.open(file_path)
    except Exception as exc:
        raise ExtractionError(f"Could not open PDF: {exc}") from exc

    pages_text = []
    try:
        for page_index in range(len(doc)):
            page = doc[page_index]
            native_text = page.get_text().strip()

            if len(native_text) >= MIN_CHARS_FOR_NATIVE_TEXT:
                pages_text.append(native_text)
                continue

            # Likely a scanned/image page — render it and OCR it.
            try:
                ocr_text = _ocr_page(page)
                pages_text.append(ocr_text)
            except Exception as exc:
                raise ExtractionError(
                    f"OCR failed on page {page_index + 1}: {exc}"
                ) from exc
    finally:
        doc.close()

    full_text = "\n\n".join(p for p in pages_text if p.strip())
    if not full_text.strip():
        raise ExtractionError("No text could be extracted (empty or unreadable document).")
    return full_text


def _ocr_page(page) -> str:
    matrix = fitz.Matrix(OCR_RENDER_ZOOM, OCR_RENDER_ZOOM)
    pixmap = page.get_pixmap(matrix=matrix)
    image = Image.open(io.BytesIO(pixmap.tobytes("png")))
    return pytesseract.image_to_string(image).strip()


def _extract_docx(file_path: str) -> str:
    try:
        doc = DocxDocument(file_path)
    except Exception as exc:
        raise ExtractionError(f"Could not open DOCX: {exc}") from exc

    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]

    # Also pull text out of tables — contracts frequently put terms in tables.
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    paragraphs.append(cell.text.strip())

    full_text = "\n\n".join(paragraphs)
    if not full_text.strip():
        raise ExtractionError("No text could be extracted (empty document).")
    return full_text
